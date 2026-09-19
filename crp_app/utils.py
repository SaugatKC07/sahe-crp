from io import BytesIO
from datetime import datetime
from django.template.loader import render_to_string
from django.conf import settings
from django.core.mail import send_mail
from django.utils import timezone

try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

from .models import Registration, Course

def generate_registration_pdf(registration):
    """Generate PDF confirmation for course registration"""
    if not PDF_AVAILABLE:
        return None
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
    
    # SAHE Brand Colors
    SAHE_GREEN = colors.HexColor('#0f7a5a')
    SAHE_LIGHT = colors.HexColor('#faf9f5')
    
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name='SAHEHeader',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=SAHE_GREEN,
        spaceAfter=30,
        alignment=TA_CENTER,
    ))
    
    styles.add(ParagraphStyle(
        name='SAHESubheader',
        parent=styles['Heading2'],
        fontSize=16,
        textColor=colors.black,
        spaceAfter=12,
    ))
    
    story = []
    
    # Header
    story.append(Paragraph("Southern Academy Higher Education", styles['SAHEHeader']))
    story.append(Spacer(1, 0.2*inch))
    
    # Registration Confirmation
    story.append(Paragraph("Course Registration Confirmation", styles['SAHESubheader']))
    story.append(Spacer(1, 0.3*inch))
    
    # Registration Details Table
    data = [
        ['Registration Details', ''],
        ['Registration ID:', str(registration.id)],
        ['Student Name:', registration.student.get_full_name() or registration.student.email],
        ['Student Email:', registration.student.email],
        ['Course Code:', registration.course.code],
        ['Course Name:', registration.course.name],
        ['Credits:', str(registration.course.credits)],
        ['Semester:', registration.semester],
        ['Registration Date:', registration.registration_date.strftime('%B %d, %Y')],
        ['Status:', registration.status.upper()],
    ]
    
    table = Table(data, colWidths=[2*inch, 4*inch])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (1, 0), SAHE_GREEN),
        ('TEXTCOLOR', (0, 0), (1, 0), colors.white),
        ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), SAHE_LIGHT),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SAHE_LIGHT]),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    
    story.append(table)
    story.append(Spacer(1, 0.5*inch))
    
    # Course Schedule Information
    schedules = registration.course.schedules.all()
    if schedules:
        story.append(Paragraph("Course Schedule", styles['SAHESubheader']))
        story.append(Spacer(1, 0.2*inch))
        
        schedule_data = [['Day', 'Time', 'Room', 'Instructor']]
        for schedule in schedules:
            schedule_data.append([
                schedule.get_day_display(),
                f"{schedule.time_slot.start_time} - {schedule.time_slot.end_time}",
                f"{schedule.room.building} - {schedule.room.name}",
                schedule.instructor.user.get_full_name() if schedule.instructor else 'TBA'
            ])
        
        schedule_table = Table(schedule_data, colWidths=[1.5*inch, 1.5*inch, 2*inch, 2*inch])
        schedule_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), SAHE_GREEN),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        
        story.append(schedule_table)
        story.append(Spacer(1, 0.5*inch))
    
    # Footer Information
    story.append(Paragraph("Important Information:", styles['SAHESubheader']))
    footer_text = """
    • Please arrive 15 minutes before the first class
    • Bring your student ID and registration confirmation
    • Contact student.services@sahe.edu.au for any queries
    • This confirmation serves as proof of registration
    """
    story.append(Paragraph(footer_text, styles['Normal']))
    story.append(Spacer(1, 0.5*inch))
    
    # Footer with date
    story.append(Paragraph(f"Generated on: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}", styles['Normal']))
    
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def export_courses_to_excel():
    """Export courses data to Excel file"""
    if not EXCEL_AVAILABLE:
        return None
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Courses"
    
    # SAHE Brand Colors
    SAHE_GREEN = '0f7a5a'
    SAHE_LIGHT = 'faf9f5'
    
    # Header style
    header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color=SAHE_GREEN, end_color=SAHE_GREEN, fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers
    headers = [
        'Course Code', 'Course Name', 'Credits', 'Level', 'Department',
        'Status', 'Capacity', 'Enrolled', 'Instructor', 'Start Date', 'End Date'
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Data rows
    courses = Course.objects.select_related('department', 'instructor__user').all()
    
    for row_num, course in enumerate(courses, 2):
        ws.cell(row=row_num, column=1, value=course.code)
        ws.cell(row=row_num, column=2, value=course.name)
        ws.cell(row=row_num, column=3, value=course.credits)
        ws.cell(row=row_num, column=4, value=course.get_level_display())
        ws.cell(row=row_num, column=5, value=course.department.name if course.department else '')
        ws.cell(row=row_num, column=6, value=course.get_status_display())
        ws.cell(row=row_num, column=7, value=course.max_capacity)
        ws.cell(row=row_num, column=8, value=course.current_enrollment)
        ws.cell(row=row_num, column=9, value=course.instructor.user.get_full_name() if course.instructor else '')
        ws.cell(row=row_num, column=10, value=course.start_date.strftime('%Y-%m-%d') if course.start_date else '')
        ws.cell(row=row_num, column=11, value=course.end_date.strftime('%Y-%m-%d') if course.end_date else '')
        
        # Apply borders to data cells
        for col_num in range(1, 12):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='left', vertical='center')
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save to buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_rows_to_excel(headers, rows, title='Report'):
    """Create a consistently styled workbook for scoped operational reports."""
    if not EXCEL_AVAILABLE:
        return b''
    wb = Workbook()
    ws = wb.active
    ws.title = (title or 'Report')[:31]
    header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0f7a5a', end_color='0f7a5a', fill_type='solid')
    border = Border(left=Side(style='thin'), right=Side(style='thin'),
                    top=Side(style='thin'), bottom=Side(style='thin'))
    for col, value in enumerate(headers, 1):
        cell = ws.cell(1, col, value)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row_num, row in enumerate(rows, 2):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row_num, col, value)
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for column in ws.columns:
        letter = column[0].column_letter
        ws.column_dimensions[letter].width = min(max(len(str(c.value or '')) for c in column) + 2, 40)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def send_registration_email(registration, status='pending', reason=''):
    """Send email notification about registration status"""
    subject = ''
    message = ''
    
    if status == 'pending':
        subject = 'Course Registration Submitted - Southern Academy Higher Education'
        message = f"""
Dear {registration.student.get_full_name() or registration.student.email},

Your registration for {registration.course.code} - {registration.course.name} has been successfully submitted.

Registration Details:
- Registration ID: {registration.id}
- Course: {registration.course.code} - {registration.course.name}
- Credits: {registration.course.credits}
- Semester: {registration.semester}
- Registration Date: {registration.registration_date.strftime('%B %d, %Y')}

Your registration is currently pending approval. You will receive another email once your registration has been reviewed.

If you have any questions, please contact student.services@sahe.edu.au.

Best regards,
Southern Academy Higher Education
"""
    elif status == 'approved':
        subject = 'Course Registration Approved - Southern Academy Higher Education'
        message = f"""
Dear {registration.student.get_full_name() or registration.student.email},

Great news! Your registration for {registration.course.code} - {registration.course.name} has been approved.

Registration Details:
- Registration ID: {registration.id}
- Course: {registration.course.code} - {registration.course.name}
- Credits: {registration.course.credits}
- Semester: {registration.semester}
- Approval Date: {registration.approval_date.strftime('%B %d, %Y')}

Please check your course schedule and ensure you attend the first class. You can download your registration confirmation from the student portal.

If you have any questions, please contact student.services@sahe.edu.au.

Best regards,
Southern Academy Higher Education
"""
    elif status == 'rejected':
        subject = 'Course Registration Update - Southern Academy Higher Education'
        message = f"""
Dear {registration.student.get_full_name() or registration.student.email},

We regret to inform you that your registration for {registration.course.code} - {registration.course.name} has been declined.

Registration Details:
- Registration ID: {registration.id}
- Course: {registration.course.code} - {registration.course.name}
- Semester: {registration.semester}

{f'Reason: {reason}' if reason else ''}

If you have any questions or would like to discuss this decision, please contact student.services@sahe.edu.au.

Best regards,
Southern Academy Higher Education
"""
    
    try:
        send_mail(
            subject,
            message,
            settings.DEFAULT_FROM_EMAIL,
            [registration.student.email],
            fail_silently=False,
        )
        return True
    except Exception as e:
        print(f"Email sending failed: {e}")
        return False

def generate_schedule_pdf(course):
    """Generate PDF for course schedule"""
    if not PDF_AVAILABLE:
        return None
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
    
    # SAHE Brand Colors
    SAHE_GREEN = colors.HexColor('#0f7a5a')
    SAHE_LIGHT = colors.HexColor('#faf9f5')
    
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name='SAHEHeader',
        parent=styles['Heading1'],
        fontSize=24,
        textColor=SAHE_GREEN,
        spaceAfter=30,
        alignment=TA_CENTER,
    ))
    
    story = []
    
    # Header
    story.append(Paragraph("Southern Academy Higher Education", styles['SAHEHeader']))
    story.append(Spacer(1, 0.2*inch))
    
    # Course Information
    story.append(Paragraph(f"Course Schedule: {course.code} - {course.name}", styles['Heading2']))
    story.append(Spacer(1, 0.3*inch))
    
    # Course Details
    course_data = [
        ['Course Code:', course.code],
        ['Course Name:', course.name],
        ['Credits:', str(course.credits)],
        ['Level:', course.get_level_display()],
        ['Instructor:', course.instructor.user.get_full_name() if course.instructor else 'TBA'],
        ['Start Date:', course.start_date.strftime('%B %d, %Y') if course.start_date else 'TBA'],
        ['End Date:', course.end_date.strftime('%B %d, %Y') if course.end_date else 'TBA'],
    ]
    
    course_table = Table(course_data, colWidths=[2*inch, 4*inch])
    course_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (1, 0), SAHE_GREEN),
        ('TEXTCOLOR', (0, 0), (1, 0), colors.white),
        ('FONTNAME', (0, 0), (1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (1, 0), 12),
        ('BOTTOMPADDING', (0, 0), (1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), SAHE_LIGHT),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, SAHE_LIGHT]),
        ('FONTNAME', (0, 1), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, -1), 10),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    
    story.append(course_table)
    story.append(Spacer(1, 0.5*inch))
    
    # Schedule Table
    schedules = course.schedules.select_related('room', 'time_slot', 'instructor').all()
    if schedules:
        story.append(Paragraph("Class Schedule", styles['Heading2']))
        story.append(Spacer(1, 0.2*inch))
        
        schedule_data = [['Day', 'Time', 'Room', 'Building', 'Instructor']]
        for schedule in schedules:
            schedule_data.append([
                schedule.get_day_display(),
                f"{schedule.time_slot.start_time} - {schedule.time_slot.end_time}",
                schedule.room.name,
                schedule.room.building,
                schedule.instructor.user.get_full_name() if schedule.instructor else 'TBA'
            ])
        
        schedule_table = Table(schedule_data, colWidths=[1.2*inch, 1.5*inch, 1.5*inch, 1.5*inch, 2*inch])
        schedule_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), SAHE_GREEN),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('GRID', (0, 0), (-1, -1), 1, colors.grey),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        
        story.append(schedule_table)
    
    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
